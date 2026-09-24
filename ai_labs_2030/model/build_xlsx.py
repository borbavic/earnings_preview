"""
Builds AI_labs_revenue_2030_bottomup.xlsx with live formulas from the assumptions in
revenue_2030_model.py. Run, then recalculate with LibreOffice (see verify_xlsx.py).

Sheets: Summary | Inputs | OpenAI | Anthropic | Deck_check | Sensitivity | Sources
Convention: blue = hardcoded input, black = formula, green = link to another sheet,
yellow fill = cells the user is meant to edit (Custom scenario column).
"""
from __future__ import annotations

import os
from typing import Callable, Dict, List

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from revenue_2030_model import ASSUMPTIONS, MIX_2026, SEGMENTS, YEARS_2026_TO_2030

FONT = "Arial"
F_BLUE = Font(name=FONT, color="0000FF")
F_BLACK = Font(name=FONT, color="000000")
F_GREEN = Font(name=FONT, color="008000")
F_BOLD = Font(name=FONT, bold=True)
F_TITLE = Font(name=FONT, bold=True, size=14)
F_HEAD = Font(name=FONT, bold=True, color="FFFFFF")
F_NOTE = Font(name=FONT, italic=True, color="595959", size=9)
FILL_HEAD = PatternFill("solid", fgColor="1F3864")
FILL_GROUP = PatternFill("solid", fgColor="D9E1F2")
FILL_YELLOW = PatternFill("solid", fgColor="FFFF00")
FILL_TOTAL = PatternFill("solid", fgColor="E2EFDA")

UNIT_FMT = {
    "%": "0.0%", "M people": "#,##0", "$k/yr": "#,##0", "$B": "#,##0.0;(#,##0.0);-", "$T": "0.0",
    "$/mo": "$#,##0", "$/yr": "$#,##0", "GW": "0.0", "1e15 tokens": "0.0", "$/M": "$0.00",
}
SCEN_COLS_IN = {"bear": "D", "base": "E", "bull": "F", "custom": "G"}      # Inputs sheet
SCEN_COLS_CO = {"bear": "C", "base": "D", "bull": "E", "custom": "F"}      # company sheets
CO_TO_IN = {"C": "D", "D": "E", "E": "F", "F": "G"}

ROW: Dict[str, int] = {}  # assumption key -> row on Inputs


def style_header(ws, row: int, cols: int):
    for c in range(1, cols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = F_HEAD
        cell.fill = FILL_HEAD
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def set_widths(ws, widths: List[float]):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


# ---------------------------------------------------------------------------
# Inputs
# ---------------------------------------------------------------------------
def build_inputs(wb: Workbook):
    ws = wb.active
    ws.title = "Inputs"
    ws["A1"] = "Inputs - 2030 assumptions (bear / base / bull) + Custom scenario"
    ws["A1"].font = F_TITLE
    ws["A2"] = ("Legend: blue = hardcoded input; yellow Custom column (G) is the one to edit - it defaults to Base. "
                "All formulas on the other sheets read this sheet. Anchors are 2026 observed data points (see Sources).")
    ws["A2"].font = F_NOTE
    headers = ["Key", "Assumption (2030 unless stated)", "Unit", "Bear", "Base", "Bull", "Custom (edit)", "2026 anchor / rationale"]
    for i, h in enumerate(headers, start=1):
        ws.cell(row=4, column=i, value=h)
    style_header(ws, 4, len(headers))
    r = 5
    group = None
    for a in ASSUMPTIONS:
        if a.group != group:
            group = a.group
            ws.cell(row=r, column=1, value=group).font = F_BOLD
            for c in range(1, 9):
                ws.cell(row=r, column=c).fill = FILL_GROUP
            r += 1
        ws.cell(row=r, column=1, value=a.key).font = F_NOTE
        ws.cell(row=r, column=2, value=a.label).font = F_BLACK
        ws.cell(row=r, column=3, value=a.unit).font = F_BLACK
        fmt = UNIT_FMT.get(a.unit, "#,##0.00")
        for col, val in zip(("D", "E", "F", "G"), (a.bear, a.base, a.bull, a.base)):
            cell = ws[f"{col}{r}"]
            cell.value = val
            cell.font = F_BLUE
            cell.number_format = fmt
        ws[f"G{r}"].fill = FILL_YELLOW
        ws.cell(row=r, column=8, value=a.anchor).font = F_NOTE
        ROW[a.key] = r
        r += 1
    ws.freeze_panes = "D5"
    set_widths(ws, [12, 52, 11, 10, 10, 10, 13, 110])
    return ws


def inp(key: str, co_col: str) -> str:
    """Reference to an Inputs cell for the scenario column co_col of a company sheet."""
    return f"Inputs!${CO_TO_IN[co_col]}${ROW[key]}"


def inp_base(key: str) -> str:
    return f"Inputs!$E${ROW[key]}"


# ---------------------------------------------------------------------------
# Company sheets
# ---------------------------------------------------------------------------
def build_company(wb: Workbook, co: str, name: str) -> Dict[str, int]:
    ws = wb.create_sheet(name)
    R: Dict[str, int] = {}
    ws["A1"] = f"{name} - 2030 revenue, bottom-up ($B unless stated)"
    ws["A1"].font = F_TITLE
    ws["A2"] = "Green = link to Inputs; black = formula; blue = hardcoded (2026 mix estimates, targets, constants)."
    ws["A2"].font = F_NOTE
    headers = ["Line", "Unit", "Bear", "Base", "Bull", "Custom", "Formula / note"]
    for i, h in enumerate(headers, start=1):
        ws.cell(row=4, column=i, value=h)
    style_header(ws, 4, len(headers))
    r = 5

    def group(title: str):
        nonlocal r
        ws.cell(row=r, column=1, value=title).font = F_BOLD
        for c in range(1, 8):
            ws.cell(row=r, column=c).fill = FILL_GROUP
        r += 1

    def line(key: str, label: str, unit: str, fn: Callable[[str], str], fmt: str, note: str = "",
             font: Font = F_BLACK, total: bool = False):
        nonlocal r
        ws.cell(row=r, column=1, value=label).font = F_BOLD if total else F_BLACK
        ws.cell(row=r, column=2, value=unit).font = F_BLACK
        for col in ("C", "D", "E", "F"):
            cell = ws[f"{col}{r}"]
            cell.value = fn(col)
            cell.font = font
            cell.number_format = fmt
            if total:
                cell.fill = FILL_TOTAL
                cell.font = Font(name=FONT, bold=True, color=font.color)
        ws.cell(row=r, column=7, value=note).font = F_NOTE
        R[key] = r
        r += 1

    B, P = "#,##0.0;(#,##0.0);-", "0.0%"

    group("Global pools (all vendors) - shared with the other company")
    line("sps_dev", "Spend per paid seat - developers", "$/yr", lambda c: f"={inp('sal_dev', c)}*{inp('ratio_dev', c)}*1000", "$#,##0", "salary x AI spend ratio", F_GREEN)
    line("sps_pro", "Spend per paid seat - professionals", "$/yr", lambda c: f"={inp('sal_pro', c)}*{inp('ratio_pro', c)}*1000", "$#,##0", "", F_GREEN)
    line("sps_gen", "Spend per paid seat - general KW", "$/yr", lambda c: f"={inp('sal_gen', c)}*{inp('ratio_gen', c)}*1000", "$#,##0", "", F_GREEN)
    line("seats_dev", "Paid developer seats", "M", lambda c: f"={inp('kw_dev', c)}*{inp('pen_dev', c)}", "#,##0.0", "workers x penetration", F_GREEN)
    line("seats_pro", "Paid professional seats", "M", lambda c: f"={inp('kw_pro', c)}*{inp('pen_pro', c)}", "#,##0.0", "", F_GREEN)
    line("seats_gen", "Paid general KW seats", "M", lambda c: f"={inp('kw_gen', c)}*{inp('pen_gen', c)}", "#,##0.0", "", F_GREEN)
    line("pool_dev", "Developer AI spend pool", "$B", lambda c: f"={c}{R['seats_dev']}*{c}{R['sps_dev']}/1000", B, "seats (M) x $/seat / 1000")
    line("pool_pro", "Professional AI spend pool", "$B", lambda c: f"={c}{R['seats_pro']}*{c}{R['sps_pro']}/1000", B, "")
    line("pool_gen", "General KW AI spend pool", "$B", lambda c: f"={c}{R['seats_gen']}*{c}{R['sps_gen']}/1000", B, "")
    line("pool_mach", "Machine / agent API pool", "$B", lambda c: f"={inp('auto_pool', c)}*{inp('auto_share', c)}*{inp('auto_capture', c)}*1000", B, "labour pool ($T) x automated share x vendor capture", F_GREEN)
    line("pool_tot", "Total enterprise AI pool (all vendors)", "$B", lambda c: f"=SUM({c}{R['pool_dev']}:{c}{R['pool_mach']})", B, "", total=True)
    r += 1

    group(f"{name} revenue 2030 by segment")
    line("seg_dev", SEGMENTS[0], "$B", lambda c: f"={c}{R['pool_dev']}*{inp('sh_dev_' + co, c)}", B, "pool x company share", F_GREEN)
    line("seg_pro", SEGMENTS[1], "$B", lambda c: f"={c}{R['pool_pro']}*{inp('sh_pro_' + co, c)}", B, "", F_GREEN)
    line("seg_gen", SEGMENTS[2], "$B", lambda c: f"={c}{R['pool_gen']}*{inp('sh_gen_' + co, c)}", B, "", F_GREEN)
    line("seg_mach", SEGMENTS[3], "$B", lambda c: f"={c}{R['pool_mach']}*{inp('sh_auto_' + co, c)}", B, "", F_GREEN)
    line("seg_subs", SEGMENTS[4], "$B", lambda c: f"={inp(co + '_mau', c)}*{inp(co + '_conv', c)}*{inp(co + '_arpu', c)}*12/1000", B, "MAU x paid conversion x ARPU x 12", F_GREEN)
    if co == "oai":
        line("seg_ads", SEGMENTS[5], "$B", lambda c: f"={inp('oai_mau', c)}*(1-{inp('oai_conv', c)})*{inp('oai_ad_arpu', c)}/1000", B, "free MAU x ads ARPU", F_GREEN)
        line("seg_comm", SEGMENTS[6], "$B", lambda c: f"={inp('oai_gmv', c)}*{inp('oai_take', c)}", B, "GMV x take rate", F_GREEN)
    else:
        line("seg_ads", SEGMENTS[5], "$B", lambda c: "=0", B, "Anthropic: no ads (stated policy) - hardcoded 0", F_BLUE)
        line("seg_comm", SEGMENTS[6], "$B", lambda c: "=0", B, "not modelled - hardcoded 0", F_BLUE)
    line("seg_other", SEGMENTS[7], "$B", lambda c: f"={inp(co + '_other', c)}", B, "", F_GREEN)
    line("total", "TOTAL REVENUE 2030", "$B", lambda c: f"=SUM({c}{R['seg_dev']}:{c}{R['seg_other']})", B, "", total=True)
    line("ent", "Enterprise subtotal (4 engines)", "$B", lambda c: f"=SUM({c}{R['seg_dev']}:{c}{R['seg_mach']})", B, "")
    line("cons", "Consumer subtotal", "$B", lambda c: f"=SUM({c}{R['seg_subs']}:{c}{R['seg_comm']})", B, "")
    line("ent_pct", "Enterprise share of revenue", "%", lambda c: f"=IF({c}{R['total']}=0,0,{c}{R['ent']}/{c}{R['total']})", P, "")
    r += 1

    group("Supply-side check (compute)")
    line("gw", "GW online (avg 2030)", "GW", lambda c: f"={inp(co + '_gw', c)}", "0.0", "", F_GREEN)
    line("rgw_inf", "Revenue per GW of inference", "$B/GW", lambda c: f"={inp('tok_per_gw', c)}*{inp('price_tok', c)}*{inp('util', c)}", B, "tokens/GW (1e15) x $/M tokens x utilisation", F_GREEN)
    line("rgw", "Revenue per GW of TOTAL capacity", "$B/GW", lambda c: f"={c}{R['rgw_inf']}*{inp('inf_share', c)}", B, "x inference share of fleet", F_GREEN)
    line("sup_cap", "Supply-side revenue capacity", "$B", lambda c: f"={c}{R['gw']}*{c}{R['rgw']}", B, "GW x $/GW", total=True)
    line("dem", "Demand-side revenue (bottom-up)", "$B", lambda c: f"={c}{R['total']}", B, "")
    line("dem_sup", "Demand / supply capacity", "x", lambda c: f"=IF({c}{R['sup_cap']}=0,0,{c}{R['dem']}/{c}{R['sup_cap']})", "0.00x", ">1.0x = compute-constrained")
    line("gw_need", "GW needed for demand at this $/GW", "GW", lambda c: f"=IF({c}{R['rgw']}=0,0,{c}{R['dem']}/{c}{R['rgw']})", "0.0", "")
    line("rgw_need", "$/GW needed for demand at this GW", "$B/GW", lambda c: f"=IF({c}{R['gw']}=0,0,{c}{R['dem']}/{c}{R['gw']})", B, "")
    line("rgwi_need", "$/GW-inference needed at this GW and inference share", "$B/GW", lambda c: f"=IF({c}{R['gw']}*{inp('inf_share', c)}=0,0,{c}{R['dem']}/({c}{R['gw']}*{inp('inf_share', c)}))", B, "demand / (GW x inference share)", F_GREEN)
    r += 1

    group("Momentum path (annualised run-rate)")
    line("rr26", "Exit-2026 run-rate", "$B", lambda c: f"={inp(co + '_rr26', c)}", B, "", F_GREEN)
    prev = "rr26"
    for y in (27, 28, 29, 30):
        key = f"rr{y}"
        line(key, f"Exit-20{y} run-rate", "$B", (lambda c, p=prev, yy=y: f"={c}{R[p]}*(1+{inp(co + '_g' + str(yy), c)})"), B, f"x (1 + growth 20{y})", F_GREEN)
        prev = key
    line("bu_vs_mom", "Bottom-up 2030 / momentum exit-2030", "x", lambda c: f"=IF({c}{R['rr30']}=0,0,{c}{R['total']}/{c}{R['rr30']})", "0.00x", "<1.0x = bottom-up is below the run-rate path")
    r += 1

    group("Bridge: 2026 run-rate mix (est., $B) -> 2030 (multiple per scenario; CAGR on Base)")
    hdr = ["Segment", "2026 RR est.", "Bear x", "Base x", "Bull x", "Custom x", "CAGR (Base, 4.33 yrs)"]
    for i, h in enumerate(hdr, start=1):
        ws.cell(row=r, column=i, value=h).font = F_BOLD
    r += 1
    seg_keys = ["seg_dev", "seg_pro", "seg_gen", "seg_mach", "seg_subs", "seg_ads", "seg_comm", "seg_other"]
    first_bridge = r
    for seg, sk in zip(SEGMENTS, seg_keys):
        ws.cell(row=r, column=1, value=seg)
        b = ws.cell(row=r, column=2, value=MIX_2026[co][seg])
        b.font = F_BLUE
        b.number_format = B
        for col in ("C", "D", "E", "F"):
            cell = ws[f"{col}{r}"]
            cell.value = f'=IF($B{r}=0,"n/a",{col}{R[sk]}/$B{r})'
            cell.number_format = '0.0"x"'
        g = ws[f"G{r}"]
        g.value = f'=IF($B{r}=0,"n/a",(D{R[sk]}/$B{r})^(1/{YEARS_2026_TO_2030})-1)'
        g.number_format = P
        r += 1
    ws.cell(row=r, column=1, value="TOTAL").font = F_BOLD
    ws[f"B{r}"] = f"=SUM(B{first_bridge}:B{r-1})"
    ws[f"B{r}"].number_format = B
    for col in ("C", "D", "E", "F"):
        ws[f"{col}{r}"] = f"={col}{R['total']}/$B{r}"
        ws[f"{col}{r}"].number_format = '0.0"x"'
    ws[f"G{r}"] = f"=(D{R['total']}/$B{r})^(1/{YEARS_2026_TO_2030})-1"
    ws[f"G{r}"].number_format = P
    R["bridge_total"] = r
    ws.cell(row=r, column=7 + 1, value="2026 mix is our estimate from disclosed run-rates (blue = editable)").font = F_NOTE
    r += 2

    group("What you need to believe (targets vs Base structure)")
    ws.cell(row=r, column=1, value="Target revenue 2030 ($B)").font = F_BOLD
    ws.cell(row=r, column=2, value="input")
    for col, v in (("C", 350), ("D", 400)):
        cell = ws[f"{col}{r}"]
        cell.value = v
        cell.font = F_BLUE
        cell.fill = FILL_YELLOW
        cell.number_format = B
    R["tgt"] = r
    r += 1

    def wline(label: str, const, fn: Callable[[str], str], fmt: str, note: str = ""):
        nonlocal r
        ws.cell(row=r, column=1, value=label)
        if const is not None:
            b = ws.cell(row=r, column=2, value=const)
            b.font = F_BLUE
        for col in ("C", "D"):
            cell = ws[f"{col}{r}"]
            cell.value = fn(col)
            cell.number_format = fmt
        ws.cell(row=r, column=7, value=note).font = F_NOTE
        r += 1
        return r - 1

    t = R["tgt"]
    wline("Implied $/GW (total) at Base GW", None, lambda c: f"={c}${t}/$D${R['gw']}", B, "target / Base GW")
    wline("Implied $/GW-inference at Base GW and Base inference share", None, lambda c: f"={c}${t}/($D${R['gw']}*{inp_base('inf_share')})", B, "target / (Base GW x Base inference share)")
    wline("Implied GW at deck economics ($/GW-inf x inference share, Inputs Base)", None, lambda c: f"={c}${t}/({inp_base('deck_rgw_inf')}*{inp_base('deck_inf_share')})", "0.0", "target / ($30B x 52.5%)")
    wline("Implied GW at Base $/GW", None, lambda c: f"={c}${t}/$D${R['rgw']}", "0.0", "")
    up = wline("Required uplift on enterprise engines vs Base", None, lambda c: f"=({c}${t}-($D${R['total']}-$D${R['ent']}))/$D${R['ent']}", "0.00x", "consumer + other held at Base")
    wline("...if only developer spend/salary moves", None, lambda c: f"={inp_base('ratio_dev')}*(1+({c}{up}-1)*$D${R['ent']}/$D${R['seg_dev']})", P, "vs Base ratio on Inputs")
    wline("...if only machine-API automated share moves", None, lambda c: f"={inp_base('auto_share')}*(1+({c}{up}-1)*$D${R['ent']}/$D${R['seg_mach']})", P, "")
    wline("...if only professional-seat penetration moves", None, lambda c: f"={inp_base('pen_pro')}*(1+({c}{up}-1)*$D${R['ent']}/$D${R['seg_pro']})", P, ">100% = impossible")
    wline("Target as % of global enterprise AI pool (Base)", None, lambda c: f"={c}${t}/$D${R['pool_tot']}", P, "")
    wline("Target as % of global labour compensation ($B) =", 65000, lambda c: f"={c}${t}/$B{r}", "0.00%", "~$58T 2025 growing to ~$65T")
    wline("Target as % of global software spend ($B) =", 2400, lambda c: f"={c}${t}/$B{r}", P, "Gartner $1.47T 2026 -> ~$2.4T 2030")

    ws.freeze_panes = "C5"
    set_widths(ws, [46, 10, 12, 12, 12, 12, 60, 40])
    return R


# ---------------------------------------------------------------------------
# Sensitivity
# ---------------------------------------------------------------------------
def build_sensitivity(wb: Workbook):
    ws = wb.create_sheet("Sensitivity")
    ws["A1"] = "Two-way sensitivities (all constants read the Base column of Inputs; headers in blue are editable)"
    ws["A1"].font = F_TITLE
    r = 3

    def grid(title, row_label, col_label, row_vals, col_vals, fn, row_fmt, col_fmt, cell_fmt):
        nonlocal r
        ws.cell(row=r, column=1, value=title).font = F_BOLD
        r += 1
        ws.cell(row=r, column=1, value=f"{row_label} \\ {col_label}").font = F_NOTE
        for j, cv in enumerate(col_vals, start=2):
            cell = ws.cell(row=r, column=j, value=cv)
            cell.font = F_BLUE
            cell.number_format = col_fmt
        hdr = r
        r += 1
        for rv in row_vals:
            cell = ws.cell(row=r, column=1, value=rv)
            cell.font = F_BLUE
            cell.number_format = row_fmt
            for j in range(2, 2 + len(col_vals)):
                cl = get_column_letter(j)
                c = ws.cell(row=r, column=j, value=fn(f"$A{r}", f"{cl}${hdr}"))
                c.number_format = cell_fmt
            r += 1
        r += 1

    grid("Revenue ($B) = total GW x deck inference share (Inputs Base) x $/GW-inference", "GW total", "$B/GW-inference",
         [10, 15, 20, 25, 30], [15, 20, 25, 30, 35, 40],
         lambda rc, cc: f"={rc}*{inp_base('deck_inf_share')}*{cc}", '0" GW"', '"$"0"B/GW-inf"', "#,##0")
    grid("Revenue ($B) at deck GW (Inputs Base) = inference share x $/GW-inference", "inference share", "$B/GW-inference",
         [0.45, 0.50, 0.55, 0.60, 0.65], [15, 20, 25, 30, 35, 40],
         lambda rc, cc: f"={inp_base('deck_gw')}*{rc}*{cc}", "0%", '"$"0"B/GW-inf"', "#,##0")
    grid("$/GW-inference ($B) = tokens/GW-yr (1e15) x realised $/M tokens x utilisation (Base)",
         "q tokens/GW-yr", "$/M tokens", [12, 16, 20, 24, 28, 32], [0.8, 1.2, 1.6, 2.0, 2.4, 3.0],
         lambda rc, cc: f"={rc}*{cc}*{inp_base('util')}", '0"q"', '"$"0.0"/M"', "#,##0.0")
    grid("OpenAI ads revenue ($B) = free MAU (M) x ads ARPU ($/yr)", "free MAU (M)", "$/free user/yr",
         [1400, 1900, 2300], [15, 30, 45, 60, 80], lambda rc, cc: f"={rc}*{cc}/1000", "#,##0", '"$"0', "#,##0")
    grid("OpenAI consumer subscriptions ($B) = MAU (M) x paid conversion x Base ARPU x 12", "MAU (M)", "paid conversion",
         [1500, 2000, 2500], [0.04, 0.06, 0.08, 0.10], lambda rc, cc: f"={rc}*{cc}*{inp_base('oai_arpu')}*12/1000", "#,##0", "0%", "#,##0")
    grid("Anthropic developer engine ($B) = paid devs (Base: workers x penetration) x spend ($k/dev/yr) x Anthropic share",
         "$k/dev/yr", "Anthropic share", [4, 6, 9, 12, 15], [0.30, 0.36, 0.42, 0.48],
         lambda rc, cc: f"={inp_base('kw_dev')}*{inp_base('pen_dev')}*{rc}*{cc}", '"$"0"k"', "0%", "#,##0")
    grid("Professional-seat global pool ($B) = professionals (Base) x penetration x $/seat", "penetration", "$/seat/yr",
         [0.30, 0.45, 0.60], [900, 1800, 2700, 3600], lambda rc, cc: f"={inp_base('kw_pro')}*{rc}*{cc}/1000", "0%", '"$"#,##0', "#,##0")
    grid("Machine / agent API global pool ($B) = labour pool (Base, $T) x automated share x vendor capture", "automated share", "vendor capture",
         [0.05, 0.10, 0.15, 0.20], [0.15, 0.20, 0.25, 0.30, 0.35], lambda rc, cc: f"={inp_base('auto_pool')}*1000*{rc}*{cc}", "0%", "0%", "#,##0")
    grid("Gross margin on inference = 1 - (compute cost per GW-yr / revenue per GW of inference)", "$B revenue/GW-inf", "$B cost/GW-yr",
         [15, 21.8, 30, 40, 50], [6, 8, 10, 13.3], lambda rc, cc: f"=1-{cc}/{rc}", '"$"0.0"B"', '"$"0.0"B"', "0%")
    set_widths(ws, [34, 12, 12, 12, 12, 12, 12])


# ---------------------------------------------------------------------------
# Deck check
# ---------------------------------------------------------------------------
def build_deck_check(wb: Workbook) -> Dict[str, int]:
    ws = wb.create_sheet("Deck_check")
    R: Dict[str, int] = {}
    ws["A1"] = "Deck check: revenue = total GW x share of capacity in inference x revenue per GW of inference"
    ws["A1"].font = F_TITLE
    ws["A2"] = "Deck inputs live on Inputs (group 'Deck'). Realised anchors apply the same 50-55% inference share to today's fleets."
    ws["A2"].font = F_NOTE
    headers = ["Line", "Unit", "Bear", "Base", "Bull", "Custom", "Formula / note"]
    for i, h in enumerate(headers, start=1):
        ws.cell(row=4, column=i, value=h)
    style_header(ws, 4, len(headers))
    r = 5
    B = "#,##0.0;(#,##0.0);-"

    def line(key, label, unit, fn, fmt, note="", font=F_BLACK, total=False):
        nonlocal r
        ws.cell(row=r, column=1, value=label).font = F_BOLD if total else F_BLACK
        ws.cell(row=r, column=2, value=unit)
        for col in ("C", "D", "E", "F"):
            cell = ws[f"{col}{r}"]
            cell.value = fn(col)
            cell.font = Font(name=FONT, bold=True, color=font.color) if total else font
            cell.number_format = fmt
            if total:
                cell.fill = FILL_TOTAL
        ws.cell(row=r, column=7, value=note).font = F_NOTE
        R[key] = r
        r += 1

    line("gw", "Deck: total GW per company", "GW", lambda c: f"={inp('deck_gw', c)}", "0.0", "", F_GREEN)
    line("sh", "Deck: share of capacity in inference", "%", lambda c: f"={inp('deck_inf_share', c)}", "0.0%", "", F_GREEN)
    line("rgwi", "Deck: revenue per GW of inference", "$B/GW", lambda c: f"={inp('deck_rgw_inf', c)}", B, "", F_GREEN)
    line("rev", "Deck revenue = GW x share x $/GW-inference", "$B", lambda c: f"={c}{R['gw']}*{c}{R['sh']}*{c}{R['rgwi']}", B, "", total=True)
    line("rgwt", "Deck implied $/GW of TOTAL capacity", "$B/GW", lambda c: f"={c}{R['sh']}*{c}{R['rgwi']}", B, "share x $/GW-inference")
    line("model", "Model 2030 $/GW-inference (tokens/GW x price x utilisation)", "$B/GW", lambda c: f"={inp('tok_per_gw', c)}*{inp('price_tok', c)}*{inp('util', c)}", B, "same scenario column on Inputs", F_GREEN)
    line("vs_model", "Deck $/GW-inference vs model", "x", lambda c: f"=IF({c}{R['model']}=0,0,{c}{R['rgwi']}/{c}{R['model']})", "0.00x", "")
    line("oai_real", "Realised OpenAI 2025 $/GW-inference", "$B/GW", lambda c: f"={inp('oai_rr_2025', c)}/({inp('oai_gw_2025', c)}*{inp('inf_share_today', c)})", B, "run-rate / (GW x inference share today)", F_GREEN)
    line("ant_real", "Realised Anthropic 2026E $/GW-inference", "$B/GW", lambda c: f"={inp('ant_rr26', c)}/({inp('ant_gw_2026', c)}*{inp('inf_share_today', c)})", B, "exit-2026 run-rate / (GW x inference share today)", F_GREEN)
    line("vs_oai", "Deck vs OpenAI 2025", "x", lambda c: f"=IF({c}{R['oai_real']}=0,0,{c}{R['rgwi']}/{c}{R['oai_real']})", "0.00x", "")
    line("vs_ant", "Deck vs Anthropic 2026E", "x", lambda c: f"=IF({c}{R['ant_real']}=0,0,{c}{R['rgwi']}/{c}{R['ant_real']})", "0.00x", "")
    r += 1
    line("gm8", "Gross margin on inference at $8B cost per GW-yr", "%", lambda c: f"=IF({c}{R['rgwi']}=0,0,1-8/{c}{R['rgwi']})", "0%", "owned / custom-silicon compute (~$8B per GW-yr)")
    line("gm13", "Gross margin on inference at $13.3B cost per GW-yr", "%", lambda c: f"=IF({c}{R['rgwi']}=0,0,1-13.3/{c}{R['rgwi']})", "0%", "rented at Oracle-OpenAI pricing ($300B/5yr/4.5GW)")
    ws.freeze_panes = "C5"
    set_widths(ws, [58, 10, 12, 12, 12, 12, 55])
    return R


# ---------------------------------------------------------------------------
# Summary + Sources
# ---------------------------------------------------------------------------
def build_summary(wb: Workbook, R: Dict[str, Dict[str, int]]):
    ws = wb.create_sheet("Summary", 0)
    ws["A1"] = "OpenAI & Anthropic - 2030 revenue framework (bottom-up vs compute vs momentum)"
    ws["A1"].font = F_TITLE
    ws["A2"] = "All cells are links to the company sheets (green). Edit assumptions on Inputs (Custom column)."
    ws["A2"].font = F_NOTE
    headers = ["Metric ($B)", "Bear", "Base", "Bull", "Custom"]
    for i, h in enumerate(headers, start=1):
        ws.cell(row=4, column=i, value=h)
    style_header(ws, 4, len(headers))
    r = 5
    rows = [
        ("OpenAI - bottom-up revenue 2030", "OpenAI", "total"),
        ("OpenAI - supply-side capacity 2030", "OpenAI", "sup_cap"),
        ("OpenAI - momentum exit-2030 run-rate", "OpenAI", "rr30"),
        ("OpenAI - $/GW needed for bottom-up", "OpenAI", "rgw_need"),
        ("Anthropic - bottom-up revenue 2030", "Anthropic", "total"),
        ("Anthropic - supply-side capacity 2030", "Anthropic", "sup_cap"),
        ("Anthropic - momentum exit-2030 run-rate", "Anthropic", "rr30"),
        ("Anthropic - $/GW needed for bottom-up", "Anthropic", "rgw_need"),
    ]
    for label, sheet, key in rows:
        ws.cell(row=r, column=1, value=label)
        for j, col in enumerate(("C", "D", "E", "F"), start=2):
            cell = ws.cell(row=r, column=j, value=f"={sheet}!{col}{R[sheet][key]}")
            cell.font = F_GREEN
            cell.number_format = "#,##0.0;(#,##0.0);-"
        r += 1
    ws.cell(row=r, column=1, value="Combined bottom-up revenue 2030").font = F_BOLD
    for j, col in enumerate(("B", "C", "D", "E"), start=2):
        cell = ws.cell(row=r, column=j, value=f"={col}5+{col}9")
        cell.number_format = "#,##0.0;(#,##0.0);-"
        cell.fill = FILL_TOTAL
        cell.font = F_BOLD
    r += 1
    for label, key, fmt in (("Deck check: 25 GW x inference share x $/GW-inference (per company)", "rev", "#,##0.0;(#,##0.0);-"),
                            ("Deck $/GW-inference vs model 2030", "vs_model", "0.00x"),
                            ("Deck $/GW-inference vs realised Anthropic 2026E", "vs_ant", "0.00x")):
        ws.cell(row=r, column=1, value=label)
        for j, col in enumerate(("C", "D", "E", "F"), start=2):
            cell = ws.cell(row=r, column=j, value=f"=Deck_check!{col}{R['Deck_check'][key]}")
            cell.font = F_GREEN
            cell.number_format = fmt
        r += 1
    r += 1
    ws.cell(row=r, column=1, value="Reference points (2026)").font = F_BOLD
    r += 1
    refs = [
        ("OpenAI run-rate Aug-2026", 40, "press reports; enterprise >50% of mix"),
        ("Anthropic run-rate end-Jul-2026", 65, "Bloomberg/Axios 17-Aug-2026"),
        ("OpenAI company plan 2030", 280, ">$280B (Bloomberg 20-Feb-2026); later reports up to ~$350B"),
        ("Anthropic company plan 2028", 195, "$190-200B (Reuters 14-Aug-2026)"),
        ("Superforecaster median, OpenAI+Anthropic combined 2030 run-rate", 300, "Forecasting Research Institute; 34% prob. >$400B"),
        ("Bain: AI revenue needed by 2030 (whole industry)", 2000, "Bain Global Technology Report 2025"),
    ]
    for label, v, note in refs:
        ws.cell(row=r, column=1, value=label)
        c = ws.cell(row=r, column=2, value=v)
        c.font = F_BLUE
        c.number_format = "#,##0"
        ws.cell(row=r, column=3, value=note).font = F_NOTE
        r += 1
    set_widths(ws, [58, 14, 14, 14, 14])


SOURCES = [
    ("OpenAI annualised revenue run-rate >$40B; enterprise >50% of mix", "Aug-2026", "press (enterprisedna.co / valueaddvc.com summaries of company statements)", "https://enterprisedna.co/resources/news/openai-enterprise-revenue-overtakes-consumer-40-billion-arr-2026/"),
    ("ChatGPT ~1B MAU (Jun-26), ~1B WAU (Jul-26); >50M consumer subs; 9M paying business users; 7M enterprise seats", "Jul/Aug-2026", "demandsage / secondtalent (secondary)", "https://www.demandsage.com/chatgpt-statistics/"),
    ("OpenAI forecasts 2030 revenue >$280B; consumer/enterprise ~50/50 by 2030", "20-Feb-2026", "Bloomberg", "https://www.bloomberg.com/news/articles/2026-02-20/openai-forecasts-its-revenue-will-top-280-billion-in-2030"),
    ("OpenAI ads plan: $2.5B 2026, $11B 2027, $25B 2028, $53B 2029, $100B 2030", "09-Apr-2026", "Axios", "https://axios.com/2026/04/09/openai-100-billion-in-ad-revenue"),
    ("ChatGPT ads: launched 9-Feb-2026 (US Free + Go); $1B run-rate in <200 days; ~$60 CPM", "2026", "press (secondary)", "https://explainx.ai/blog/openai-advertise-in-chatgpt-ads-launch-july-2026"),
    ("OpenAI API 15B tokens/min; Codex 2M+ weekly users", "Mar-2026", "OpenAI", "https://openai.com/index/a-business-that-scales-with-the-value-of-intelligence/"),
    ("OpenAI compute: 0.2GW (2023) -> 1.9GW (2025) while run-rate $2B -> $20B (~$10B/GW)", "2025", "MBI Deep Dives citing OpenAI", "https://www.mbi-deepdives.com/oai-breakeven/"),
    ("OpenAI 30GW target by 2030; compute spend ~$600B -> ~$750B through 2030", "2026", "WSJ / DCD (secondary)", "https://www.datacenterdynamics.com/en/news/openai-cuts-projected-compute-spend-to-600bn-by-2030-report/"),
    ("OpenAI-Nvidia 10GW LOI; AMD 6GW; Broadcom 10GW; Stargate 10GW (8GW+ secured)", "2025-26", "OpenAI", "https://openai.com/index/openai-nvidia-systems-partnership/"),
    ("Oracle-OpenAI $300B / 5 yrs / 4.5GW (~$13.3B per GW-year)", "Jul/Sep-2025", "The Register", "https://www.theregister.com/2025/07/22/openai_oracle_gpus/"),
    ("Anthropic run-rate $65B (end-Jul-26); $47B mid-May; $9B end-2025", "17-Aug-2026", "Bloomberg / TechCrunch / Axios", "https://techcrunch.com/2026/08/17/anthropics-annualized-revenue-surges-to-65b/"),
    ("Anthropic run-rate >$30B; >1,000 customers >$1M/yr; multi-GW TPU capacity from 2027", "Apr-2026", "Anthropic", "https://www.anthropic.com/news/google-broadcom-partnership-compute"),
    ("Anthropic IPO case: $190-200B revenue in 2028", "14-Aug-2026", "Reuters", "https://www.usnews.com/news/top-news/articles/2026-08-14/exclusive-anthropic-ipo-valuation-hinges-on-190-200-billion-2028-revenue-forecast-sources-say"),
    ("Anthropic ~5GW end-2026, ~10GW 2027; 3.5GW Google/Broadcom TPU; AWS up to 5GW Trainium; 2.16GW Australia", "2026", "press (secondary)", "https://cryptobriefing.com/anthropic-10-gw-compute-capacity-2027/"),
    ("Claude Code $8B ARR (May-26); >50% enterprise", "May-2026", "press (secondary)", "https://aibusinessweekly.net/p/claude-code-statistics"),
    ("Anthropic gross margin 40% (2025), guided 63% (2026), 70% (2027)", "Jan-2026", "The Information via valueaddvc (secondary)", "https://valueaddvc.com/blog/is-anthropic-profitable-2026-losses-burn-rate-and-the-path-to-breakeven"),
    ("Claude Code cost: avg ~$6/dev/day; 90% of users <$12/day; enterprise avg $150-250/dev/month", "2026", "Anthropic docs / morphllm", "https://code.claude.com/docs/en/costs"),
    ("Claude pricing: Pro $20, Max $100/$200, Team $25 std / $125 premium, Enterprise ~$20/seat + usage", "Sep-2026", "Anthropic", "https://claude.com/pricing"),
    ("ChatGPT Go $8/mo US, Rs399 India (free promo to Dec-26); Plus $20; Pro $200", "2026", "OpenAI", "https://openai.com/index/introducing-chatgpt-go/"),
    ("Meta FY2025 revenue $200.1B; DAP 3.58B; Q4-25 ARPP $16.56 (~$58/yr)", "28-Jan-2026", "Meta IR", "https://investor.atmeta.com/investor-news/press-release-details/2026/Meta-Reports-Fourth-Quarter-and-Full-Year-2025-Results/default.aspx"),
    ("Google Search & other ads revenue 2025: $175.8B", "Feb-2026", "Alphabet Q4-25 release", "https://s206.q4cdn.com/479360582/files/doc_financials/2025/q4/2025q4-alphabet-earnings-release.pdf"),
    ("Google processes 3.2 quadrillion tokens/month (7x YoY)", "May-2026", "Google I/O 2026", "https://blog.google/innovation-and-ai/sundar-pichai-io-2026/"),
    ("M365 Copilot 20M paid seats (Q3 FY26) at $30/seat; 450M M365 commercial seats; GitHub Copilot 4.7M paid", "Jan-Apr 2026", "Microsoft earnings via Directions on Microsoft (secondary)", "https://www.directionsonmicrosoft.com/microsoft-claims-15-million-paid-m365-copilot-seats/"),
    ("Cursor $2B+ ARR", "2026", "secondtalent (secondary)", "https://www.secondtalent.com/resources/cursor-vs-github-copilot/"),
    ("Enterprise LLM API spend share: Anthropic ~32-40%, OpenAI ~25-27%; coding ~54% Anthropic", "2025-26", "Menlo Ventures (via secondary)", "https://medium.com/@david.j.sea/anthropic-just-passed-openai-in-revenue-here-is-why-it-matters-e3dd9bb04069"),
    ("AI spend per employee 2026: avg $2,068 (+50% YoY); top decile >$2,800; median <$200", "06-May-2026", "Atlanta Fed", "https://www.atlantafed.org/research-and-data/publications/policy-hub-macroblog/2026/05/06/how-much-firms-spending-on-ai-and-what-will-happen-to-headcounts"),
    ("Gartner: IT spend 2026 $6.37T (+14.2%); software $1.468T (+15.5%); AI spending $2.59T (+47%)", "Jul-2026", "Gartner", "https://www.gartner.com/en/newsroom/press-releases/2026-07-27-gartner-forecasts-worldwide-it-spending-to-grow-14-point-2-percent-in-2026-totaling-6-point-37-trillion"),
    ("Knowledge workers ~1B+ globally (Gartner); 1.25B information workers (Forrester 2018); global labour comp ~$58T", "various", "Forrester / Gartner / Substrate gist (secondary)", "https://go.forrester.com/blogs/the-global-information-worker-population-swells-to-1-25-billion-in-2018/"),
    ("Bain: AI needs $2T annual revenue by 2030; $800B shortfall; ~200GW global AI compute by 2030", "Sep-2025", "Bain & Company", "https://www.bain.com/about/media-center/press-releases/20252/$2-trillion-in-new-revenue-needed-to-fund-ais-scaling-trend---bain--companys-6th-annual-global-technology-report/"),
    ("1GW AI data center TCO: ~$38B upfront capex + ~$0.9B/yr opex (servers ~60% of annualised cost)", "2026", "Epoch AI", "https://epoch.ai/data-insights/ai-datacenter-cost-breakdown"),
    ("Frontier API list prices Sep-2026: Claude Opus 5.5 $4/$20, GPT-6 Sol $2/$10, Gemini 3.8 Flash $0.75/$3.75 per M tokens", "Sep-2026", "benchlm / developersdigest (secondary)", "https://benchlm.ai/llm-pricing"),
    ("Superforecasters: OpenAI+Anthropic combined 2030 run-rate median $300B; 34% prob >$400B; 18% prob <$100B", "2026", "Forecasting Research Institute", "https://forecastingresearch.substack.com/p/will-the-ai-boom-continue-forecasting"),
]


def build_sources(wb: Workbook):
    ws = wb.create_sheet("Sources")
    ws["A1"] = "Data anchors and sources (as of 24-Sep-2026; several are press summaries of company disclosures - flagged 'secondary')"
    ws["A1"].font = F_TITLE
    headers = ["Data point", "Date", "Source", "URL"]
    for i, h in enumerate(headers, start=1):
        ws.cell(row=3, column=i, value=h)
    style_header(ws, 3, len(headers))
    for i, (dp, dt, src, url) in enumerate(SOURCES, start=4):
        ws.cell(row=i, column=1, value=dp)
        ws.cell(row=i, column=2, value=dt)
        ws.cell(row=i, column=3, value=src)
        c = ws.cell(row=i, column=4, value=url)
        c.hyperlink = url
        c.font = Font(name=FONT, color="0563C1", underline="single")
    set_widths(ws, [95, 14, 45, 90])


def main(out: str):
    wb = Workbook()
    build_inputs(wb)
    R = {"OpenAI": build_company(wb, "oai", "OpenAI"), "Anthropic": build_company(wb, "ant", "Anthropic")}
    R["Deck_check"] = build_deck_check(wb)
    build_sensitivity(wb)
    build_sources(wb)
    build_summary(wb, R)
    # default font everywhere it was not set explicitly
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is not None and (cell.font is None or cell.font.name != FONT):
                    f = cell.font
                    cell.font = Font(name=FONT, bold=f.bold, italic=f.italic, color=f.color, size=f.size, underline=f.underline)
    wb.save(out)
    print("saved", out)


if __name__ == "__main__":
    here = os.path.dirname(os.path.abspath(__file__))
    main(os.path.join(here, "AI_labs_revenue_2030_bottomup.xlsx"))
