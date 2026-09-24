"""Cross-checks the recalculated workbook against revenue_2030_model.py and the v3 reference (run after recalc)."""
import os
import sys

from openpyxl import load_workbook

from revenue_2030_model import SCEN, SEGMENTS, company, deck_check, momentum, supply, v3, yearly

here = os.path.dirname(os.path.abspath(__file__))
path = sys.argv[1] if len(sys.argv) > 1 else os.path.join(here, "AI_labs_revenue_2030_bottomup.xlsx")
wb = load_workbook(path, data_only=True)
cols = {"bear": "C", "base": "D", "bull": "E", "custom": "F"}
bad = 0


def rows_of(ws):
    out = {}
    for rr in range(1, ws.max_row + 1):
        v = ws.cell(row=rr, column=1).value
        if v and v not in out:  # first occurrence (segment names repeat in the bridge block)
            out[v] = rr
    return out


def close(a, b, tol=1e-6):
    return a is not None and abs(a - b) <= tol * max(1.0, abs(b))


for co, sheet in (("oai", "OpenAI"), ("ant", "Anthropic")):
    ws = wb[sheet]
    rows = rows_of(ws)
    for s, col in cols.items():
        ps = "base" if s == "custom" else s
        py = company(co, ps)
        for seg in SEGMENTS + ["TOTAL"]:
            label = "TOTAL REVENUE 2030" if seg == "TOTAL" else seg
            xl = ws[f"{col}{rows[label]}"].value
            if not close(xl, py[seg]):
                bad += 1
                print(f"MISMATCH {sheet} {s} {seg}: xlsx={xl} py={py[seg]}")
        sp = supply(co, ps)
        for label, key in (("Supply-side revenue = inference GW x yield", "Supply-side revenue = inf GW x yield ($B)"),
                           ("Gross margin (on supply-side revenue)", "Gross margin (supply-side revenue)"),
                           ("Yield needed for the bottom-up at this GW and share", "Yield needed for bottom-up at this GW and share ($B)")):
            xl = ws[f"{col}{rows[label]}"].value
            if not close(xl, sp[key]):
                bad += 1
                print(f"MISMATCH {sheet} {s} {label}: xlsx={xl} py={sp[key]}")
        mo = momentum(co, ps)
        for label, key in (("Exit-2030 run-rate", "Exit-2030 run-rate"), ("Calendar revenue 2030", "Calendar revenue 2030 (avg of exit run-rates)")):
            xl = ws[f"{col}{rows[label]}"].value
            if not close(xl, mo[key]):
                bad += 1
                print(f"MISMATCH {sheet} {s} {label}: xlsx={xl} py={mo[key]}")
    # Base must equal v3 exactly
    for label, key in (("TOTAL REVENUE 2030", "rev_2030"), ("Supply-side revenue = inference GW x yield", "rev_2030"),
                       ("Calendar revenue 2030", "rev_2030"), ("Exit-2030 run-rate", "exit_arr_2030"),
                       ("Gross margin (on supply-side revenue)", "gross_margin_2030"),
                       ("Revenue per average total GW-year", "rev_per_total_gw_2030")):
        xl = ws[f"D{rows[label]}"].value
        if not close(xl, v3(co, key), 1e-9):
            bad += 1
            print(f"V3 MISMATCH {sheet} {label}: xlsx={xl} v3={v3(co, key)}")
    print(f"{sheet}: base total xlsx={ws['D' + str(rows['TOTAL REVENUE 2030'])].value:.6f}  v3={v3(co, 'rev_2030'):.6f}  "
          f"GM xlsx={ws['D' + str(rows['Gross margin (on supply-side revenue)'])].value:.4%} v3={v3(co, 'gross_margin_2030'):.4%}")

def check_yearly(wb_, scen, tol=1e-6):
    """Compare the year-by-year block (columns B..F) of both company sheets with yearly(co, scen)."""
    n_bad, n_ok = 0, 0
    for co, sheet in (("oai", "OpenAI"), ("ant", "Anthropic")):
        ws_ = wb_[sheet]
        rows_ = rows_of(ws_)
        py = yearly(co, scen)
        for label, vals in py.items():
            if label not in rows_:
                n_bad += 1
                print(f"YEARLY label missing on {sheet}: {label}")
                continue
            rr = rows_[label]
            for j, v in enumerate(vals):
                xl = ws_.cell(row=rr, column=2 + j).value
                if xl is None or abs(xl - v) > tol * max(1.0, abs(v)):
                    n_bad += 1
                    print(f"YEARLY MISMATCH {sheet} [{scen}] {label} {2026 + j}: xlsx={xl} py={v}")
                else:
                    n_ok += 1
        # in Base the supply rows must equal the v3 yearly series
        if scen == "base":
            for label, key in (("Calendar revenue path ($B)", "rev"), ("Supply: average total capacity (GW)", "gw_avg_total"),
                               ("Supply: revenue per average inference GW-year ($B)", "yield_inf"), ("Supply: gross margin", "gross_margin"),
                               ("Momentum: exit run-rate ($B)", "exit_arr")):
                for j, v in enumerate(v3(co, "rev_2030") and __import__("revenue_2030_model").v3s(co, key)):
                    xl = ws_.cell(row=rows_[label], column=2 + j).value
                    if abs(xl - v) > 1e-9 * max(1.0, abs(v)):
                        n_bad += 1
                        print(f"YEARLY vs v3 MISMATCH {sheet} {label} {2026 + j}: xlsx={xl} v3={v}")
    print(f"yearly [{scen}]: {n_ok} cells ok, {n_bad} bad")
    return n_bad


bad += check_yearly(wb, "base")

# selector test: switch both sheets to Bull, recalculate a scratch copy, compare with yearly(co, 'bull')
import shutil, subprocess, tempfile
from openpyxl import load_workbook as _lw
scratch = os.path.join(tempfile.gettempdir(), "yearly_selector_test.xlsx")
wb_f = _lw(path)  # formulas
for sheet in ("OpenAI", "Anthropic"):
    ws_f = wb_f[sheet]
    for rr in range(1, ws_f.max_row + 1):
        lab = ws_f.cell(row=rr, column=1).value
        if isinstance(lab, str) and lab.startswith("Year-by-year evolution"):
            ws_f.cell(row=rr, column=3).value = "Bull"
wb_f.save(scratch)
recalc = os.path.join(os.path.dirname(os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))), "")  # placeholder
RECALC = os.environ.get("RECALC_PY", "")
if RECALC and os.path.exists(RECALC):
    res = subprocess.run([sys.executable, RECALC, scratch, "240"], capture_output=True, text=True)
    print("selector test recalc:", res.stdout.strip().replace("\n", " ")[:120])
    bad += check_yearly(_lw(scratch, data_only=True), "bull")
else:
    print("selector test skipped (set RECALC_PY to the recalc.py path to run it)")

ws = wb["Deck_check"]
drows = rows_of(ws)
for s, col in cols.items():
    dc = deck_check("base" if s == "custom" else s)
    xl = ws[f"{col}{drows['Deck revenue = GW x share x $/GW-inference']}"].value
    if not close(xl, dc["Deck revenue = GW x inf share x $/GW-inf ($B)"]):
        bad += 1
        print("MISMATCH deck revenue", s, xl)
print("Deck_check base revenue:", ws[f"D{drows['Deck revenue = GW x share x $/GW-inference']}"].value)

ws = wb["Summary"]
srows = rows_of(ws)
print("Summary OpenAI bottom-up:", [ws.cell(row=srows['OpenAI - bottom-up revenue 2030'], column=c).value for c in range(2, 6)])
print("Summary Anthropic bottom-up:", [ws.cell(row=srows['Anthropic - bottom-up revenue 2030'], column=c).value for c in range(2, 6)])
print("Summary combined:", [ws.cell(row=srows['Combined bottom-up revenue 2030'], column=c).value for c in range(2, 6)])
# reconciliation block: every difference must be ~0
n_recon = 0
for rr in range(1, ws.max_row + 1):
    lab = ws.cell(row=rr, column=1).value
    if isinstance(lab, str) and (lab.startswith("OpenAI: ") or lab.startswith("Anthropic: ")):
        d = ws.cell(row=rr, column=4).value
        n_recon += 1
        if d is None or abs(d) > 1e-9:
            bad += 1
            print("RECON MISMATCH", lab, ws.cell(row=rr, column=2).value, ws.cell(row=rr, column=3).value, d)
print(f"Reconciliation lines checked: {n_recon}")
print("MISMATCHES:", bad)
sys.exit(1 if bad else 0)
